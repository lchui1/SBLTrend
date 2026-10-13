"""
sbl_engine.py

Core SBL (Statistical Bin Limit) Trend Processing & Visualization Engine
for the SBLTrend Web Application.
"""

import io
import os
import re
import tempfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend for web servers
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


# ---------------------------------------------------------------------
# COLUMN SPECIFICATION & CANDIDATES
# ---------------------------------------------------------------------
SPECIFIED_FULL_DATA_COLUMNS = [
    "SITE",
    "LOT",
    "PRODUCT",
    "OPER",
    "PROC",
    "SCREENPLAYNAME",
    "SCREENNAME",
    "SCREENENDDATE",
    "WAFERLIST",
    "CONDITIONNAME",
    "CONDITIONLIMIT",
    "ACTUALVALUE",
    "LOTWFROP",
    "PASS",
]

CONDITION_COL_CANDIDATES = [
    "CONDITIONNAME",
    "CONDITION_NAME",
    "CONDITION",
    "COND_NAME",
    "TRIGGER_CONDITION",
    "STATUS",
    "RESULT",
    "DISPOSITION",
    "SBL_RESULT",
]

TIME_COL_CANDIDATES = [
    "SCREENENDDATE",
    "SCREEN_END_DATE",
    "WW",
    "WORKWEEK",
    "WORK_WEEK",
    "YEAR_WW",
    "EVENT_DATE",
    "DATE",
    "TIMESTAMP",
    "DATETIME",
    "TEST_DATE",
    "MODIFIED_DATE",
]

LOT_COL_CANDIDATES = [
    "LOT",
    "TEST_LOT",
    "SORT_LOT",
    "LOT_ID",
    "LOT_NAME",
    "VISUAL_ID",
]

SITE_COL_CANDIDATES = [
    "SITE",
    "SCREEN_FACILITY",
    "FACILITY",
    "FAB",
    "PLANT",
    "LOCATION",
]

PRODUCT_COL_CANDIDATES = [
    "PRODUCT",
    "PRODUCT_NAME",
    "DEVICE",
    "PROD",
    "PROGRAM",
]

SCREENPLAY_COL_CANDIDATES = [
    "SCREENPLAYNAME",
    "SCREENPLAY_NAME",
    "SCREENPLAY",
    "SCRIPT_NAME",
]

SCREEN_COL_CANDIDATES = [
    "SCREENNAME",
    "SCREEN_NAME",
    "SCREEN",
]


# ---------------------------------------------------------------------
# HELPER UTILITIES
# ---------------------------------------------------------------------
def find_best_match(columns: List[str], candidates: List[str]) -> Optional[str]:
    """Find the best matching column name in a case-insensitive manner."""
    col_map = {str(c).strip().upper(): c for c in columns}
    for candidate in candidates:
        cand_upper = candidate.upper()
        if cand_upper in col_map:
            return col_map[cand_upper]

    for candidate in candidates:
        cand_upper = candidate.upper().replace("_", "")
        for col_upper, orig_col in col_map.items():
            if cand_upper == col_upper.replace("_", ""):
                return orig_col

    for candidate in candidates:
        cand_upper = candidate.upper()
        for col_upper, orig_col in col_map.items():
            if cand_upper in col_upper or col_upper in cand_upper:
                return orig_col
    return None


def detect_delimiter_from_bytes(content: bytes) -> str:
    """Detect CSV delimiter (comma, tab, semicolon, pipe) from byte sample."""
    sample = content[:4096].decode("utf-8", errors="ignore")
    delimiters = [",", "\t", ";", "|"]
    counts = {d: sample.count(d) for d in delimiters}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def parse_intel_ww(series: pd.Series) -> pd.Series:
    """Standardize Intel Work Week format (e.g. 'WW2432', '2432')."""
    def _format_ww(val):
        if pd.isna(val):
            return "Unknown"
        s = str(val).strip().upper()
        m = re.search(r"(?:20)?(\d{2})[-_ ]*(?:WW)?[-_ ]*(\d{2})", s)
        if m:
            yy, ww = m.group(1), m.group(2)
            return f"WW{yy}{ww.zfill(2)}"
        return s

    return series.apply(_format_ww)


def extract_workweek(df: pd.DataFrame, time_col: Optional[str]) -> Tuple[pd.DataFrame, str]:
    """Derive standardized Work Week (WW<YY><WW>) from date/time or existing WW column."""
    df = df.copy()
    if not time_col:
        time_col = find_best_match(list(df.columns), TIME_COL_CANDIDATES)

    ww_label = "WorkWeek"

    if time_col and time_col in df.columns:
        col_data = df[time_col]
        sample_str = str(col_data.dropna().iloc[0]) if not col_data.dropna().empty else ""
        is_ww_str = bool(re.search(r"WW\d{2}|\b\d{4}\b", sample_str, re.IGNORECASE)) and "WW" in time_col.upper()

        if is_ww_str:
            df[ww_label] = parse_intel_ww(col_data)
        else:
            parsed_dates = pd.to_datetime(col_data, errors="coerce")
            if parsed_dates.notna().sum() > 0:
                iso_cal = parsed_dates.dt.isocalendar()
                df[ww_label] = iso_cal.apply(
                    lambda row: f"WW{str(row['year'])[-2:]}{str(row['week']).zfill(2)}" if pd.notna(row["week"]) else "Unknown",
                    axis=1,
                )
            else:
                df[ww_label] = parse_intel_ww(col_data)
    else:
        df[ww_label] = "WW_Current"

    df[ww_label] = df[ww_label].fillna("Unknown")
    return df, ww_label


def deduplicate_sbl_data(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, str]]:
    """
    Deduplicate so that each lot is only counted once per (SCREENPLAYNAME + SCREENNAME),
    keeping the row with the oldest SCREENENDDATE (earliest chronological timestamp).
    """
    df_work = df.copy()
    col_mapping = {}

    lot_col = find_best_match(list(df_work.columns), LOT_COL_CANDIDATES)
    screenplay_col = find_best_match(list(df_work.columns), SCREENPLAY_COL_CANDIDATES)
    screen_col = find_best_match(list(df_work.columns), SCREEN_COL_CANDIDATES)
    date_col = find_best_match(list(df_work.columns), TIME_COL_CANDIDATES)

    if lot_col:
        col_mapping["LOT"] = lot_col
    if screenplay_col:
        col_mapping["SCREENPLAYNAME"] = screenplay_col
    if screen_col:
        col_mapping["SCREENNAME"] = screen_col
    if date_col:
        col_mapping["SCREENENDDATE"] = date_col

    dedup_subset = [c for c in [lot_col, screenplay_col, screen_col] if c is not None and c in df_work.columns]

    if len(dedup_subset) >= 2:
        if date_col and date_col in df_work.columns:
            df_work["_temp_sort_date"] = pd.to_datetime(df_work[date_col], errors="coerce")
            df_work = df_work.sort_values(by="_temp_sort_date", ascending=True, na_position="last")
            df_work = df_work.drop_duplicates(subset=dedup_subset, keep="first")
            df_work = df_work.drop(columns=["_temp_sort_date"])
        else:
            df_work = df_work.drop_duplicates(subset=dedup_subset, keep="first")

    return df_work, col_mapping


def extract_specified_columns(df: pd.DataFrame) -> pd.DataFrame:
    """
    Extract and standardize only the 14 requested columns:
    SITE, LOT, PRODUCT, OPER, PROC, SCREENPLAYNAME, SCREENNAME, SCREENENDDATE,
    WAFERLIST, CONDITIONNAME, CONDITIONLIMIT, ACTUALVALUE, LOTWFROP, PASS
    """
    col_dict = {}
    df_cols_lookup = {str(c).strip().upper().replace("_", ""): c for c in df.columns}

    for target_col in SPECIFIED_FULL_DATA_COLUMNS:
        norm_target = target_col.upper().replace("_", "")
        if norm_target in df_cols_lookup:
            actual_col = df_cols_lookup[norm_target]
            col_dict[target_col] = df[actual_col]
        else:
            matched = find_best_match(
                list(df.columns),
                [target_col, target_col.replace("NAME", "_NAME"), target_col.replace("DATE", "_DATE")],
            )
            if matched and matched in df.columns:
                col_dict[target_col] = df[matched]
            else:
                col_dict[target_col] = ""

    return pd.DataFrame(col_dict)


def filter_data_by_months(df: pd.DataFrame, max_months: Optional[int] = 6, time_col: Optional[str] = None) -> Tuple[pd.DataFrame, bool, Optional[str], Optional[str]]:
    """
    If dataset spans > max_months of history (default 6 months),
    limit the dataset to the most recent max_months (counting backward from the latest date in data).
    """
    if max_months is None or max_months <= 0:
        return df, False, None, None

    if not time_col:
        time_col = find_best_match(list(df.columns), TIME_COL_CANDIDATES)

    if not time_col or time_col not in df.columns:
        return df, False, None, None

    parsed_dates = pd.to_datetime(df[time_col], errors="coerce")
    valid_dates = parsed_dates.dropna()
    if valid_dates.empty:
        return df, False, None, None

    max_date = valid_dates.max()
    min_date = valid_dates.min()
    cutoff_date = max_date - pd.DateOffset(months=max_months)

    if min_date < cutoff_date:
        mask = (parsed_dates >= cutoff_date) | parsed_dates.isna()
        df_filtered = df[mask].copy()
        return df_filtered, True, cutoff_date.strftime("%Y-%m-%d"), max_date.strftime("%Y-%m-%d")
    else:
        return df, False, min_date.strftime("%Y-%m-%d"), max_date.strftime("%Y-%m-%d")


# ---------------------------------------------------------------------
# CHART GENERATORS (For Excel Embedding)
# ---------------------------------------------------------------------
def setup_plot_style():
    plt.style.use("seaborn-v0_8-whitegrid" if "seaborn-v0_8-whitegrid" in plt.style.available else "default")
    plt.rcParams["font.sans-serif"] = ["Segoe UI", "Arial", "DejaVu Sans", "Helvetica"]
    plt.rcParams["axes.edgecolor"] = "#cccccc"
    plt.rcParams["axes.linewidth"] = 0.8
    plt.rcParams["grid.color"] = "#e5e5e5"
    plt.rcParams["grid.linestyle"] = "--"
    plt.rcParams["grid.alpha"] = 0.7


def generate_ww_counts_chart(trend_df: pd.DataFrame, ww_col: str) -> str:
    setup_plot_style()
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=140)

    time_periods = trend_df[ww_col].astype(str).tolist()
    counts = trend_df["TRIGGERED_LOTS"].values
    x_indices = np.arange(len(time_periods))

    bars = ax.bar(x_indices, counts, color="#1f77b4", alpha=0.7, edgecolor="#0d47a1", width=0.55, label="Triggered Lots")
    for bar, count in zip(bars, counts):
        if count > 0:
            ax.annotate(
                f"{int(count):,}",
                xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()),
                xytext=(0, 3),
                textcoords="offset points",
                ha="center",
                va="bottom",
                fontsize=8,
                fontweight="bold",
                color="#1a237e",
            )

    ax.plot(x_indices, counts, color="#0d47a1", linewidth=1.5, marker="o", markersize=4, alpha=0.8)

    if len(counts) >= 3:
        rolling_avg = pd.Series(counts).rolling(window=3, min_periods=1).mean().to_numpy()
        ax.plot(x_indices, rolling_avg, color="#d32f2f", linewidth=2.0, linestyle="--", label="3-WW Moving Avg")

    ax.set_title("SBL Triggered Lots Trend by Work Week", fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel("Work Week", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_ylabel("Triggered Lots Count", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_xticks(x_indices)
    ax.set_xticklabels(time_periods, rotation=45, ha="right", fontsize=9)
    ax.legend(loc="upper left", frameon=True, facecolor="white")
    ax.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))

    plt.tight_layout()
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.savefig(temp_file.name, dpi=160, bbox_inches="tight")
    plt.close()
    return temp_file.name


def generate_ww_rate_chart(trend_df: pd.DataFrame, ww_col: str) -> str:
    setup_plot_style()
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=140)

    time_periods = trend_df[ww_col].astype(str).tolist()
    rates = trend_df["TRIGGER_RATE_PCT"].values
    x_indices = np.arange(len(time_periods))

    ax.plot(x_indices, rates, color="#c2185b", linewidth=2.2, marker="s", markersize=5, label="Trigger Rate (%)")
    for x, y in zip(x_indices, rates):
        ax.annotate(
            f"{y:.1f}%",
            xy=(x, y),
            xytext=(0, 5),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=8,
            fontweight="bold",
            color="#880e4f",
        )

    if len(rates) >= 3:
        rolling_avg = pd.Series(rates).rolling(window=3, min_periods=1).mean().to_numpy()
        ax.plot(x_indices, rolling_avg, color="#f57c00", linewidth=1.8, linestyle="--", label="3-WW Moving Avg (%)")

    ax.set_title("SBL Trigger Rate Trend (%) by Work Week", fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel("Work Week", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_ylabel("Trigger Rate (%)", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_xticks(x_indices)
    ax.set_xticklabels(time_periods, rotation=45, ha="right", fontsize=9)
    ax.set_ylim(0, max(10, min(100, rates.max() * 1.3) if len(rates) > 0 and rates.max() > 0 else 10))
    ax.yaxis.set_major_formatter(ticker.PercentFormatter())
    ax.legend(loc="upper left", frameon=True, facecolor="white")

    plt.tight_layout()
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.savefig(temp_file.name, dpi=160, bbox_inches="tight")
    plt.close()
    return temp_file.name


def generate_multiline_chart(pivot_df: pd.DataFrame, category_name: str, title: str, ylabel: str = "Trigger Count", is_percentage: bool = False) -> str:
    setup_plot_style()
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=140)

    time_periods = pivot_df.index.astype(str).tolist()
    x_indices = np.arange(len(time_periods))
    colors = plt.cm.tab10.colors

    for i, col in enumerate(pivot_df.columns):
        color = colors[i % len(colors)]
        ax.plot(
            x_indices,
            pivot_df[col].values,
            marker="o",
            markersize=4.5,
            linewidth=1.8,
            color=color,
            label=str(col),
        )

    ax.set_title(title, fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel("Work Week", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_ylabel(ylabel, fontsize=10, fontweight="bold", labelpad=8)
    ax.set_xticks(x_indices)
    ax.set_xticklabels(time_periods, rotation=45, ha="right", fontsize=9)
    ax.legend(
        title=category_name,
        bbox_to_anchor=(1.02, 1),
        loc="upper left",
        borderaxespad=0,
        frameon=True,
        facecolor="white",
        fontsize=8,
    )
    if is_percentage:
        ax.yaxis.set_major_formatter(ticker.PercentFormatter())
        max_val = pivot_df.max().max() if not pivot_df.empty else 10
        ax.set_ylim(0, max(10, min(100, max_val * 1.25) if max_val > 0 else 10))
    else:
        ax.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))

    plt.tight_layout()
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.savefig(temp_file.name, dpi=160, bbox_inches="tight")
    plt.close()
    return temp_file.name


def generate_stacked_chart(pivot_df: pd.DataFrame, category_name: str, title: str) -> str:
    setup_plot_style()
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=140)

    time_periods = pivot_df.index.astype(str).tolist()
    x_indices = np.arange(len(time_periods))
    colors = plt.cm.tab10.colors
    bottom = np.zeros(len(time_periods))

    for i, col in enumerate(pivot_df.columns):
        values = pivot_df[col].values
        color = colors[i % len(colors)]
        ax.bar(
            x_indices,
            values,
            bottom=bottom,
            label=str(col),
            color=color,
            edgecolor="white",
            linewidth=0.5,
            width=0.55,
        )
        bottom += values

    ax.set_title(title, fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel("Work Week", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_ylabel("Total Triggers", fontsize=10, fontweight="bold", labelpad=8)
    ax.set_xticks(x_indices)
    ax.set_xticklabels(time_periods, rotation=45, ha="right", fontsize=9)
    ax.legend(
        title=category_name,
        bbox_to_anchor=(1.02, 1),
        loc="upper left",
        borderaxespad=0,
        frameon=True,
        facecolor="white",
        fontsize=8,
    )
    ax.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))

    plt.tight_layout()
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.savefig(temp_file.name, dpi=160, bbox_inches="tight")
    plt.close()
    return temp_file.name


def generate_pareto_chart(pareto_df: pd.DataFrame, category_col: str, count_col: str, title: str, top_n: int = 12) -> str:
    setup_plot_style()
    fig, ax1 = plt.subplots(figsize=(10, 4.8), dpi=140)

    df_plot = pareto_df.head(top_n).copy()
    categories = df_plot[category_col].astype(str).tolist()
    counts = df_plot[count_col].values
    cum_pct = df_plot["CUM_PCT"].values
    x_indices = np.arange(len(categories))

    bars = ax1.bar(x_indices, counts, color="#2b5c8f", width=0.55, edgecolor="#1a365d", label="Trigger Count")
    ax1.set_xlabel(category_col, fontsize=10, fontweight="bold", labelpad=8)
    ax1.set_ylabel("Trigger Count", fontsize=10, fontweight="bold", color="#1a365d", labelpad=8)
    ax1.set_xticks(x_indices)
    ax1.set_xticklabels(categories, rotation=45, ha="right", fontsize=8.5)
    ax1.yaxis.set_major_locator(ticker.MaxNLocator(integer=True))

    for bar, count in zip(bars, counts):
        ax1.annotate(
            f"{int(count):,}",
            xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()),
            xytext=(0, 2),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=7.5,
            fontweight="bold",
        )

    ax2 = ax1.twinx()
    ax2.plot(x_indices, cum_pct, color="#e53935", marker="D", markersize=4, linewidth=1.8, label="Cumulative %")
    ax2.axhline(80, color="#fb8c00", linestyle="--", linewidth=1.1, alpha=0.8, label="80% Cutoff")
    ax2.set_ylabel("Cumulative %", fontsize=10, fontweight="bold", color="#e53935", labelpad=8)
    ax2.set_ylim(0, 105)
    ax2.yaxis.set_major_formatter(ticker.PercentFormatter())
    ax2.grid(False)

    plt.title(title, fontsize=13, fontweight="bold", pad=12)
    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc="center right", frameon=True, facecolor="white", fontsize=8)

    plt.tight_layout()
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.savefig(temp_file.name, dpi=160, bbox_inches="tight")
    plt.close()
    return temp_file.name


# ---------------------------------------------------------------------
# EXCEL FORMATTING UTILITIES
# ---------------------------------------------------------------------
def style_worksheet(ws: openpyxl.worksheet.worksheet.Worksheet, start_row: int = 1, start_col: int = 1, end_row: Optional[int] = None, end_col: Optional[int] = None):
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    data_font = Font(name="Segoe UI", size=9.5)
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    max_r = end_row or ws.max_row
    max_c = end_col or ws.max_column

    for col_idx in range(start_col, max_c + 1):
        cell = ws.cell(row=start_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row_idx in range(start_row + 1, max_r + 1):
        use_alt = (row_idx % 2 == 0)
        for col_idx in range(start_col, max_c + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill

    for col_idx in range(start_col, max_c + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(start_row, max_r + 1):
            val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def write_df_to_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, df: pd.DataFrame, start_row: int = 1, start_col: int = 1, section_title: Optional[str] = None) -> int:
    cur_row = start_row
    if section_title:
        title_cell = ws.cell(row=cur_row, column=start_col, value=section_title)
        title_cell.font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
        cur_row += 1

    for c_idx, col_name in enumerate(df.columns, start=start_col):
        ws.cell(row=cur_row, column=c_idx, value=str(col_name))

    for r_idx, row_values in enumerate(df.itertuples(index=False), start=cur_row + 1):
        for c_idx, val in enumerate(row_values, start=start_col):
            cell_val = "" if pd.isna(val) else val
            ws.cell(row=r_idx, column=c_idx, value=cell_val)

    style_worksheet(ws, start_row=cur_row, start_col=start_col, end_row=cur_row + len(df), end_col=start_col + len(df.columns) - 1)
    return cur_row + len(df) + 2


# ---------------------------------------------------------------------
# CORE ANALYSIS PIPELINE (Used by FastAPI backend)
# ---------------------------------------------------------------------
def process_sbl_data(
    file_bytes: bytes,
    filename: str,
    months: int = 6,
    filter_col: Optional[str] = None,
    filter_val: Optional[str] = None,
) -> Tuple[Dict[str, Any], bytes]:
    """
    Main processing entry point.
    Returns:
      1. analysis_result: Dict with summary metrics and chart datasets for web UI.
      2. excel_bytes: The complete multi-worksheet Excel report binary with embedded charts.
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext in [".xlsx", ".xls"]:
        df_raw = pd.read_excel(io.BytesIO(file_bytes))
    else:
        sep = detect_delimiter_from_bytes(file_bytes)
        df_raw = pd.read_csv(io.BytesIO(file_bytes), sep=sep, low_memory=False)

    df_raw.columns = [str(c).strip() for c in df_raw.columns]

    # Optional Filtering
    if filter_col and filter_val:
        if filter_col in df_raw.columns:
            val_list = [v.strip() for v in filter_val.split(",")]
            df_raw = df_raw[df_raw[filter_col].astype(str).isin(val_list)]

    # 1. 6-Month Data Window Rule
    date_col = find_best_match(list(df_raw.columns), TIME_COL_CANDIDATES) or "SCREENENDDATE"
    df_raw, is_month_filtered, min_date_str, max_date_str = filter_data_by_months(df_raw, max_months=months, time_col=date_col)

    # 2. Extract Full Data (Original within months window)
    df_full_data_original = extract_specified_columns(df_raw)

    # 3. Deduplication: 1 record per (LOT + SCREENPLAYNAME + SCREENNAME), oldest SCREENENDDATE
    df_dedup, dedup_col_map = deduplicate_sbl_data(df_raw)

    lot_col = find_best_match(list(df_dedup.columns), LOT_COL_CANDIDATES) or "LOT"
    site_col = find_best_match(list(df_dedup.columns), SITE_COL_CANDIDATES) or "SITE"
    product_col = find_best_match(list(df_dedup.columns), PRODUCT_COL_CANDIDATES) or "PRODUCT"
    screenplay_col = find_best_match(list(df_dedup.columns), SCREENPLAY_COL_CANDIDATES) or "SCREENPLAYNAME"
    condition_col = find_best_match(list(df_dedup.columns), CONDITION_COL_CANDIDATES) or "CONDITIONNAME"

    if condition_col in df_dedup.columns:
        df_dedup["IS_SBL_TRIGGER"] = (
            df_dedup[condition_col].astype(str).str.upper().str.contains("FAILS", na=False).astype(int)
        )
    else:
        df_dedup["IS_SBL_TRIGGER"] = 1

    df_dedup, ww_col = extract_workweek(df_dedup, date_col)

    # A. WW Summary
    lot_ww_summary = (
        df_dedup.groupby([ww_col, lot_col], observed=False)["IS_SBL_TRIGGER"]
        .max()
        .reset_index(name="LOT_TRIGGERED")
    )
    ww_trend_table = (
        lot_ww_summary.groupby(ww_col, observed=False)
        .agg(
            TOTAL_LOTS=(lot_col, "nunique"),
            TRIGGERED_LOTS=("LOT_TRIGGERED", "sum"),
        )
        .reset_index()
    )
    ww_trend_table["TRIGGER_RATE_PCT"] = (
        (ww_trend_table["TRIGGERED_LOTS"] / ww_trend_table["TOTAL_LOTS"]) * 100
    ).fillna(0).round(2)
    ww_trend_table = ww_trend_table.sort_values(by=ww_col).reset_index(drop=True)

    df_triggered_only = df_dedup[df_dedup["IS_SBL_TRIGGER"] == 1].copy()

    # B. Site Summary
    has_site = site_col in df_dedup.columns and len(df_dedup) > 0
    if has_site:
        site_lot_summary = df_dedup.groupby([site_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        site_summary_table = (
            site_lot_summary.groupby(site_col, observed=False)
            .agg(TOTAL_TESTED_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum"))
            .reset_index()
            .sort_values(by="TRIGGERED_LOTS", ascending=False)
            .reset_index(drop=True)
        )
        site_summary_table["TRIGGER_RATE_PCT"] = (
            (site_summary_table["TRIGGERED_LOTS"] / site_summary_table["TOTAL_TESTED_LOTS"]) * 100
        ).fillna(0).round(2)
        tot_s = site_summary_table["TRIGGERED_LOTS"].sum()
        site_summary_table["PERCENT_OF_FAILURES"] = ((site_summary_table["TRIGGERED_LOTS"] / tot_s) * 100).round(2) if tot_s > 0 else 0
        site_summary_table["CUM_PCT"] = site_summary_table["PERCENT_OF_FAILURES"].cumsum().round(2)

        top_sites = site_summary_table[site_col].head(8).tolist()
        site_ww_lot = df_dedup.groupby([ww_col, site_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        site_ww_agg = site_ww_lot.groupby([ww_col, site_col], observed=False).agg(TOTAL_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum")).reset_index()
        site_ww_agg["TRIGGER_RATE_PCT"] = ((site_ww_agg["TRIGGERED_LOTS"] / site_ww_agg["TOTAL_LOTS"]) * 100).fillna(0).round(2)
        site_ww_top = site_ww_agg[site_ww_agg[site_col].isin(top_sites)]

        site_counts_pivot = site_ww_top.pivot_table(index=ww_col, columns=site_col, values="TRIGGERED_LOTS", aggfunc="sum", fill_value=0).reindex(columns=[s for s in top_sites if s in site_ww_top[site_col].values]).sort_index()
        site_rates_pivot = site_ww_top.pivot_table(index=ww_col, columns=site_col, values="TRIGGER_RATE_PCT", aggfunc="mean", fill_value=0).reindex(columns=[s for s in top_sites if s in site_ww_top[site_col].values]).sort_index().round(2)
    else:
        site_summary_table = pd.DataFrame()
        site_counts_pivot = pd.DataFrame()
        site_rates_pivot = pd.DataFrame()

    # C. Product Summary
    has_product = product_col in df_dedup.columns and len(df_dedup) > 0
    if has_product:
        product_lot_summary = df_dedup.groupby([product_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        product_summary_table = (
            product_lot_summary.groupby(product_col, observed=False)
            .agg(TOTAL_TESTED_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum"))
            .reset_index()
            .sort_values(by="TRIGGERED_LOTS", ascending=False)
            .reset_index(drop=True)
        )
        product_summary_table["TRIGGER_RATE_PCT"] = (
            (product_summary_table["TRIGGERED_LOTS"] / product_summary_table["TOTAL_TESTED_LOTS"]) * 100
        ).fillna(0).round(2)
        tot_p = product_summary_table["TRIGGERED_LOTS"].sum()
        product_summary_table["PERCENT_OF_FAILURES"] = ((product_summary_table["TRIGGERED_LOTS"] / tot_p) * 100).round(2) if tot_p > 0 else 0
        product_summary_table["CUM_PCT"] = product_summary_table["PERCENT_OF_FAILURES"].cumsum().round(2)

        top_products = product_summary_table[product_col].head(8).tolist()
        prod_ww_lot = df_dedup.groupby([ww_col, product_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        prod_ww_agg = prod_ww_lot.groupby([ww_col, product_col], observed=False).agg(TOTAL_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum")).reset_index()
        prod_ww_agg["TRIGGER_RATE_PCT"] = ((prod_ww_agg["TRIGGERED_LOTS"] / prod_ww_agg["TOTAL_LOTS"]) * 100).fillna(0).round(2)
        prod_ww_top = prod_ww_agg[prod_ww_agg[product_col].isin(top_products)]

        product_counts_pivot = prod_ww_top.pivot_table(index=ww_col, columns=product_col, values="TRIGGERED_LOTS", aggfunc="sum", fill_value=0).reindex(columns=[p for p in top_products if p in prod_ww_top[product_col].values]).sort_index()
        product_rates_pivot = prod_ww_top.pivot_table(index=ww_col, columns=product_col, values="TRIGGER_RATE_PCT", aggfunc="mean", fill_value=0).reindex(columns=[p for p in top_products if p in prod_ww_top[product_col].values]).sort_index().round(2)
    else:
        product_summary_table = pd.DataFrame()
        product_counts_pivot = pd.DataFrame()
        product_rates_pivot = pd.DataFrame()

    # D. Screenplay Summary
    has_screenplay = screenplay_col in df_dedup.columns and len(df_dedup) > 0
    if has_screenplay:
        sp_lot_summary = df_dedup.groupby([screenplay_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        screenplay_summary_table = (
            sp_lot_summary.groupby(screenplay_col, observed=False)
            .agg(TOTAL_TESTED_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum"))
            .reset_index()
            .sort_values(by="TRIGGERED_LOTS", ascending=False)
            .reset_index(drop=True)
        )
        screenplay_summary_table["TRIGGER_RATE_PCT"] = (
            (screenplay_summary_table["TRIGGERED_LOTS"] / screenplay_summary_table["TOTAL_TESTED_LOTS"]) * 100
        ).fillna(0).round(2)
        tot_sp = screenplay_summary_table["TRIGGERED_LOTS"].sum()
        screenplay_summary_table["PERCENT_OF_FAILURES"] = ((screenplay_summary_table["TRIGGERED_LOTS"] / tot_sp) * 100).round(2) if tot_sp > 0 else 0
        screenplay_summary_table["CUM_PCT"] = screenplay_summary_table["PERCENT_OF_FAILURES"].cumsum().round(2)

        top_screenplays = screenplay_summary_table[screenplay_col].head(8).tolist()
        sp_ww_lot = df_dedup.groupby([ww_col, screenplay_col, lot_col], observed=False)["IS_SBL_TRIGGER"].max().reset_index(name="LOT_TRIGGERED")
        sp_ww_agg = sp_ww_lot.groupby([ww_col, screenplay_col], observed=False).agg(TOTAL_LOTS=(lot_col, "nunique"), TRIGGERED_LOTS=("LOT_TRIGGERED", "sum")).reset_index()
        sp_ww_agg["TRIGGER_RATE_PCT"] = ((sp_ww_agg["TRIGGERED_LOTS"] / sp_ww_agg["TOTAL_LOTS"]) * 100).fillna(0).round(2)
        sp_ww_top = sp_ww_agg[sp_ww_agg[screenplay_col].isin(top_screenplays)]

        screenplay_counts_pivot = sp_ww_top.pivot_table(index=ww_col, columns=screenplay_col, values="TRIGGERED_LOTS", aggfunc="sum", fill_value=0).reindex(columns=[s for s in top_screenplays if s in sp_ww_top[screenplay_col].values]).sort_index()
        screenplay_rates_pivot = sp_ww_top.pivot_table(index=ww_col, columns=screenplay_col, values="TRIGGER_RATE_PCT", aggfunc="mean", fill_value=0).reindex(columns=[s for s in top_screenplays if s in sp_ww_top[screenplay_col].values]).sort_index().round(2)
    else:
        screenplay_summary_table = pd.DataFrame()
        screenplay_counts_pivot = pd.DataFrame()
        screenplay_rates_pivot = pd.DataFrame()

    df_triggered_details = extract_specified_columns(df_triggered_only)

    # -----------------------------------------------------------------
    # BUILD EXCEL WORKBOOK IN MEMORY
    # -----------------------------------------------------------------
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    temp_chart_files: List[str] = []

    # 1. WW_Trend
    ws_ww = wb.create_sheet(title="WW_Trend")
    ws_ww.views.sheetView[0].showGridLines = True
    write_df_to_sheet(ws_ww, ww_trend_table, start_row=1, start_col=1, section_title="Work Week Trigger Trend Summary")
    chart_ww_counts = generate_ww_counts_chart(ww_trend_table, ww_col)
    chart_ww_rate = generate_ww_rate_chart(ww_trend_table, ww_col)
    temp_chart_files.extend([chart_ww_counts, chart_ww_rate])
    img1 = OpenpyxlImage(chart_ww_counts); img1.width, img1.height = 680, 320; ws_ww.add_image(img1, "F2")
    img2 = OpenpyxlImage(chart_ww_rate); img2.width, img2.height = 680, 320; ws_ww.add_image(img2, "F19")

    # 2. Site_Trend
    ws_site = wb.create_sheet(title="Site_Trend")
    ws_site.views.sheetView[0].showGridLines = True
    next_r_site = write_df_to_sheet(ws_site, site_summary_table, start_row=1, start_col=1, section_title="Site/Facility SBL Summary & Trigger Rate (%)")
    if not site_counts_pivot.empty:
        next_r_site = write_df_to_sheet(ws_site, site_counts_pivot.reset_index(), start_row=next_r_site, start_col=1, section_title="Triggered Lots Counts by Site (Work Week)")
        if not site_rates_pivot.empty:
            write_df_to_sheet(ws_site, site_rates_pivot.reset_index(), start_row=next_r_site, start_col=1, section_title="Trigger Rate (%) by Site (Work Week)")
        chart_s_lines = generate_multiline_chart(site_counts_pivot, "Site", "SBL Triggered Lots Trend by Site/Facility")
        chart_s_rates = generate_multiline_chart(site_rates_pivot, "Site", "SBL Trigger Rate Trend (%) by Site/Facility", ylabel="Trigger Rate (%)", is_percentage=True)
        chart_s_stacked = generate_stacked_chart(site_counts_pivot, "Site", "SBL Trigger Composition by Site Over Time")
        chart_s_pareto = generate_pareto_chart(site_summary_table, site_col, "TRIGGERED_LOTS", "Site/Facility SBL Pareto Breakdown")
        temp_chart_files.extend([chart_s_lines, chart_s_rates, chart_s_stacked, chart_s_pareto])
        img_s1 = OpenpyxlImage(chart_s_lines); img_s1.width, img_s1.height = 680, 320; ws_site.add_image(img_s1, "H2")
        img_s2 = OpenpyxlImage(chart_s_rates); img_s2.width, img_s2.height = 680, 320; ws_site.add_image(img_s2, "H19")
        img_s3 = OpenpyxlImage(chart_s_stacked); img_s3.width, img_s3.height = 680, 320; ws_site.add_image(img_s3, "H36")
        img_s4 = OpenpyxlImage(chart_s_pareto); img_s4.width, img_s4.height = 680, 320; ws_site.add_image(img_s4, "H53")

    # 3. Product_Trend
    ws_prod = wb.create_sheet(title="Product_Trend")
    ws_prod.views.sheetView[0].showGridLines = True
    next_r = write_df_to_sheet(ws_prod, product_summary_table, start_row=1, start_col=1, section_title="Product SBL Summary & Trigger Rate (%)")
    if not product_counts_pivot.empty:
        next_r = write_df_to_sheet(ws_prod, product_counts_pivot.reset_index(), start_row=next_r, start_col=1, section_title="Triggered Lots Counts by Product (Work Week)")
        if not product_rates_pivot.empty:
            write_df_to_sheet(ws_prod, product_rates_pivot.reset_index(), start_row=next_r, start_col=1, section_title="Trigger Rate (%) by Product (Work Week)")
        chart_p_lines = generate_multiline_chart(product_counts_pivot, "Product", "SBL Triggered Lots Trend by Product (Top Products)")
        chart_p_rates = generate_multiline_chart(product_rates_pivot, "Product", "SBL Trigger Rate Trend (%) by Product", ylabel="Trigger Rate (%)", is_percentage=True)
        chart_p_stacked = generate_stacked_chart(product_counts_pivot, "Product", "SBL Trigger Composition by Product Over Time")
        chart_p_pareto = generate_pareto_chart(product_summary_table, product_col, "TRIGGERED_LOTS", "Product SBL Pareto Breakdown")
        temp_chart_files.extend([chart_p_lines, chart_p_rates, chart_p_stacked, chart_p_pareto])
        img_p1 = OpenpyxlImage(chart_p_lines); img_p1.width, img_p1.height = 680, 320; ws_prod.add_image(img_p1, "H2")
        img_p2 = OpenpyxlImage(chart_p_rates); img_p2.width, img_p2.height = 680, 320; ws_prod.add_image(img_p2, "H19")
        img_p3 = OpenpyxlImage(chart_p_stacked); img_p3.width, img_p3.height = 680, 320; ws_prod.add_image(img_p3, "H36")
        img_p4 = OpenpyxlImage(chart_p_pareto); img_p4.width, img_p4.height = 680, 320; ws_prod.add_image(img_p4, "H53")

    # 4. Screenplay_Trend
    ws_sp = wb.create_sheet(title="Screenplay_Trend")
    ws_sp.views.sheetView[0].showGridLines = True
    next_r_sp = write_df_to_sheet(ws_sp, screenplay_summary_table, start_row=1, start_col=1, section_title="Screenplay SBL Summary & Trigger Rate (%)")
    if not screenplay_counts_pivot.empty:
        next_r_sp = write_df_to_sheet(ws_sp, screenplay_counts_pivot.reset_index(), start_row=next_r_sp, start_col=1, section_title="Triggered Lots Counts by Screenplay (Work Week)")
        if not screenplay_rates_pivot.empty:
            write_df_to_sheet(ws_sp, screenplay_rates_pivot.reset_index(), start_row=next_r_sp, start_col=1, section_title="Trigger Rate (%) by Screenplay (Work Week)")
        chart_sp_lines = generate_multiline_chart(screenplay_counts_pivot, "Screenplay", "SBL Triggered Lots Trend by Screenplay (Top Screenplays)")
        chart_sp_rates = generate_multiline_chart(screenplay_rates_pivot, "Screenplay", "SBL Trigger Rate Trend (%) by Screenplay", ylabel="Trigger Rate (%)", is_percentage=True)
        chart_sp_pareto = generate_pareto_chart(screenplay_summary_table, screenplay_col, "TRIGGERED_LOTS", "Screenplay SBL Pareto Breakdown")
        temp_chart_files.extend([chart_sp_lines, chart_sp_rates, chart_sp_pareto])
        img_sp1 = OpenpyxlImage(chart_sp_lines); img_sp1.width, img_sp1.height = 680, 320; ws_sp.add_image(img_sp1, "H2")
        img_sp2 = OpenpyxlImage(chart_sp_rates); img_sp2.width, img_sp2.height = 680, 320; ws_sp.add_image(img_sp2, "H19")
        img_sp3 = OpenpyxlImage(chart_sp_pareto); img_sp3.width, img_sp3.height = 680, 320; ws_sp.add_image(img_sp3, "H36")

    # 5. Triggered_Details
    ws_trig = wb.create_sheet(title="Triggered_Details")
    ws_trig.views.sheetView[0].showGridLines = True
    write_df_to_sheet(ws_trig, df_triggered_details, start_row=1, start_col=1, section_title="SBL Triggered Records (CONDITIONNAME == 'FAILS')")

    # 6. Full_Data
    ws_full = wb.create_sheet(title="Full_Data")
    ws_full.views.sheetView[0].showGridLines = True
    write_df_to_sheet(ws_full, df_full_data_original, start_row=1, start_col=1, section_title="Full Raw Data (Original Dataset - 14 Schema Columns)")

    # Save to in-memory bytes
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_bytes = excel_buffer.getvalue()

    for tmp_file in temp_chart_files:
        try:
            if os.path.exists(tmp_file):
                os.remove(tmp_file)
        except Exception:
            pass

    # -----------------------------------------------------------------
    # STRUCTURE JSON PAYLOAD FOR WEB UI
    # -----------------------------------------------------------------
    total_tested_lots = int(df_dedup[lot_col].nunique()) if lot_col in df_dedup.columns else len(df_dedup)
    total_triggered_lots = int(ww_trend_table["TRIGGERED_LOTS"].sum()) if not ww_trend_table.empty else 0
    overall_trigger_rate = round((total_triggered_lots / total_tested_lots * 100), 2) if total_tested_lots > 0 else 0

    web_payload = {
        "status": "success",
        "filename": filename,
        "summary": {
            "total_raw_rows": len(df_raw),
            "total_dedup_rows": len(df_dedup),
            "total_tested_lots": total_tested_lots,
            "total_triggered_lots": total_triggered_lots,
            "overall_trigger_rate_pct": overall_trigger_rate,
            "months_window": months,
            "date_range": f"{min_date_str} to {max_date_str}" if min_date_str else "N/A",
            "is_month_filtered": is_month_filtered,
        },
        "ww_trend": ww_trend_table.to_dict(orient="records"),
        "site_summary": site_summary_table.head(15).to_dict(orient="records") if not site_summary_table.empty else [],
        "product_summary": product_summary_table.head(15).to_dict(orient="records") if not product_summary_table.empty else [],
        "screenplay_summary": screenplay_summary_table.head(15).to_dict(orient="records") if not screenplay_summary_table.empty else [],
        "site_pivot": site_counts_pivot.reset_index().to_dict(orient="records") if not site_counts_pivot.empty else [],
        "site_rates_pivot": site_rates_pivot.reset_index().to_dict(orient="records") if not site_rates_pivot.empty else [],
        "product_pivot": product_counts_pivot.reset_index().to_dict(orient="records") if not product_counts_pivot.empty else [],
        "product_rates_pivot": product_rates_pivot.reset_index().to_dict(orient="records") if not product_rates_pivot.empty else [],
        "screenplay_pivot": screenplay_counts_pivot.reset_index().to_dict(orient="records") if not screenplay_counts_pivot.empty else [],
        "screenplay_rates_pivot": screenplay_rates_pivot.reset_index().to_dict(orient="records") if not screenplay_rates_pivot.empty else [],
        "triggered_details_sample": df_triggered_details.head(100).fillna("").to_dict(orient="records"),
        "full_data_sample": df_full_data_original.head(100).fillna("").to_dict(orient="records"),
    }

    return web_payload, excel_bytes
